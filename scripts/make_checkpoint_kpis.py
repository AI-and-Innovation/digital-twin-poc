"""AE-IT widget cards for one day as an Excel workbook (and CSVs), from the same records the Kafka feeds carry.

Runs the Hyderabad simulation for one local day and keeps the kpi.checkpoint, kpi.curbside, kpi.overview and
flight.status records (src/airport_eye/generators/kpi.py; the cards and desks are chosen in the ``kpi`` section
of config/passenger_tracking.json). Every number in the workbook is a field of one of those records.

  python scripts/make_checkpoint_kpis.py                            # 9 Oct 2026, share/hyd_checkpoint_kpis/
  python scripts/make_checkpoint_kpis.py --date 2026-10-10 --scenario peak

Needs openpyxl for the workbook (pip install -e '.[kpi]'); without it only the CSVs are written.
All passengers, flights and vehicles are synthetic.
"""

from __future__ import annotations

import argparse
import csv
import json
import sys
import zipfile
from datetime import datetime, timedelta, timezone
from pathlib import Path
from zoneinfo import ZoneInfo

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from airport_eye.esri import ZoneLocator, to_esri_record  # noqa: E402
from airport_eye.generators.passenger import TerminalSimulation, load_tracking  # noqa: E402
from airport_eye.queue_files import to_csv  # noqa: E402
from airport_eye.tracking import KPI_FEEDS, STEMS, columns_for  # noqa: E402

TZ = ZoneInfo("Asia/Kolkata")
DROP = ("event_id", "schema_version", "airport_id", "terminal_id", "source_system", "sequence_no", "ingestion_time",
        "ingestion_time_epoch_ms", "web_mercator_x", "web_mercator_y", "spatial_reference_wkid", "lat", "lon",
        "window_start", "reference_elevation_m", "global_elevation_m")


def simulate(day: datetime, scenario: str | None, step_s: float) -> tuple[list[dict], dict]:
    """Every KPI record of the day, as the flat Esri records the topics carry."""
    config, layout, inventory, _ = load_tracking(ROOT / "config/passenger_tracking.json", scenario)
    if config.kpi is None:
        raise SystemExit("config/passenger_tracking.json has no 'kpi' section")
    locator = ZoneLocator.from_layout(layout)
    start, end = day.astimezone(timezone.utc), (day + timedelta(days=1)).astimezone(timezone.utc)
    sim = TerminalSimulation(config, layout, inventory, start, end)
    out: list[dict] = []
    now = sim.warm_start
    while now < end:   # the same stepping as generate_ticks, without building the position records
        to_minute = 60 - (now.timestamp() % 60)
        step = min(step_s, to_minute) if to_minute > 1e-6 else step_s
        now += timedelta(seconds=step)
        ts = now.timestamp()
        sim.tick(ts, step)
        if sim.ground is not None:
            sim.ground.finished_landside()
        if abs(ts - round(ts)) < 1e-3 and round(ts) % 60 == 0:
            reporting = now > start
            for r in sim.minute_records(ts, now, reporting):
                if reporting and r.event_type in KPI_FEEDS:
                    out.append(to_esri_record(r, locator))
            if not reporting and sim.ground is not None:
                sim.ground.reset_window()
    return out, {"kpi": config.kpi}


def local(ms: int) -> datetime:
    return datetime.fromtimestamp(ms / 1000, TZ)


def mmss(seconds) -> str | None:
    if seconds is None:
        return None
    s = int(round(seconds))
    return f"{s // 60:02d}m {s % 60:02d}s"


def card_row(r: dict) -> dict:
    row = {"time_local": local(r["event_time_epoch_ms"]).strftime("%Y-%m-%d %H:%M")}
    for k, v in r.items():
        if k in DROP or k.startswith("today_hourly") or k in ("event_type", "event_time", "event_time_epoch_ms",
                                                                "entity_id", "window_end"):
            continue
        row[k] = v
        if k in ("avg_processing_time_seconds", "avg_wait_time_seconds", "avg_dwell_seconds"):
            row[k.replace("_seconds", "")] = mmss(v)
    return row


def write_csv(path: Path, rows: list[dict]) -> None:
    with path.open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=list(dict.fromkeys(k for r in rows for k in r)))
        w.writeheader()
        w.writerows(rows)


def formulas(kind: str, kpi) -> list[tuple[str, str]]:
    if kind == "kpi.checkpoint":
        return [
            ("Queue Length", "Passengers waiting in front of the desk now. The simulation has one queue per area in "
                             "front of all open desks; each desk is shown its share of it."),
            ("Average Processing Time", f"Mean service time of passengers this desk finished in the last "
                                        f"{kpi.rolling_window_minutes:g} minutes."),
            ("Average Wait Time", f"Mean time from joining the queue to reaching this desk, for passengers whose "
                                  f"service started in the last {kpi.rolling_window_minutes:g} minutes."),
            ("Happiness Index", f"From the average wait: Happy < {kpi.wait_threshold_minutes / 2:g} min, Neutral up to "
                                f"{kpi.wait_threshold_minutes:g} min, Unhappy above (until 'FRT' is defined)."),
            ("Wait Time Trends", f"Hourly mean wait at this desk against the {kpi.wait_threshold_minutes:g} min "
                                 "threshold."),
        ]
    if kind == "kpi.curbside":
        t = kpi.trolleys
        return [
            ("People Count", f"People within {kpi.kerb_people_radius_m:g} m of the kerb line(s) now."),
            ("Live Vehicle count", "Vehicles stopped at the kerb bays now."),
            ("Vehicles This hour", "Vehicles that stopped since the start of the clock hour. Trend % = vs the previous "
                                   "hour up to the same minute."),
            ("Avg Dwell time", "Mean stop time of vehicles that left the kerb this hour. Trend % = vs the previous hour."),
            ("Vehicles Type", "This hour's vehicles by type: private car, app cab, taxi, bus (no bikes in the "
                              "simulation)."),
            ("Trolley Availability", f"Departures only: available / {t.total}. ASSUMED model: passengers dropped at "
                                     f"the kerb take one ({t.take_share.get('domestic', 0):.0%} domestic, "
                                     f"{t.take_share.get('international', 0):.0%} international); staff collect used "
                                     f"ones every {t.collection_every_minutes:g} min."),
        ]
    if kind == "kpi.overview":
        return [
            ("Arrival / Departure Flights", "Flights that arrived (in-block) / departed (off-block) this clock hour so "
                                            "far. Trend % = vs flights scheduled (STA / STD) in the same window."),
            ("Stand Occupancy", "Stands with an aircraft or a turnaround in progress now / total contact stands."),
            ("Avg Taxi in / out Time", f"Airport card: mean AIBT - ALDT / ATOT - AOBT this hour, vs target "
                                       f"({kpi.taxi_in_target_minutes:g} / {kpi.taxi_out_target_minutes:g} min, assumed)."),
            ("Arrival / Departure Bags", "Terminal card: bags on flights that arrived / departed this hour, vs the bags "
                                         "expected from the flights scheduled this hour (threshold)."),
            ("OTA / OTD", f"% of this hour's flights within +/- {kpi.otp_window_minutes:g} min of schedule. "
                          "ATA = in-block, ATD = off-block."),
            ("Avg Delay", "Mean ATA - STA / ATD - STD in minutes (negative = early)."),
            ("PAX details", "Passengers forecast (flights scheduled that hour) and actual (flights that arrived / "
                            "departed that hour)."),
            ("Data", "In-block / off-block times, passengers and stands come from the simulation; schedules, estimates, "
                     "taxi times, forecasts and bags are SYNTHETIC per flight (Flights sheet)."),
        ]
    return []


CARD_NAMES = {"check_in": "Check-in", "security": "Security", "passport": "Passport"}


def sheet_name(r: dict) -> str:
    if r["event_type"] == "kpi.checkpoint":
        return f"{CARD_NAMES[r['card']]} {r['desk_number']:02d}"
    if r["event_type"] == "kpi.curbside":
        return f"Curbside {r['side']}"
    return "Airport overview" if r["scope"] == "airport" else "Terminal overview"


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--date", default="2026-10-09", help="local (IST) day to simulate")
    ap.add_argument("--scenario", help="named scenario from the tracking config, e.g. peak")
    ap.add_argument("--step", type=float, default=5.0, help="simulation step in seconds (5 = live)")
    ap.add_argument("--snapshot", default="auto",
                    help="local times (HH:MM, comma-separated) for the 'Cards at' sheets; auto = the busiest "
                         "check-in minute and the busiest passport minute of the day")
    ap.add_argument("--out", default=str(ROOT / "share/hyd_checkpoint_kpis"))
    ap.add_argument("--zip", default=str(ROOT / "share/AirportEye_HYD_checkpoint_KPIs.zip"))
    args = ap.parse_args()
    day = datetime.fromisoformat(args.date).replace(tzinfo=TZ)
    out = Path(args.out)
    out.mkdir(parents=True, exist_ok=True)
    print(f"Simulating Hyderabad T1, {args.date} 00:00-24:00 IST (step {args.step:g} s)...", flush=True)
    records, info = simulate(day, args.scenario, args.step)
    kpi = info["kpi"]

    cards: dict[str, list[dict]] = {}
    flights: dict[str, dict] = {}
    for r in records:
        if r["event_type"] == "flight.status":
            flights[r["flight_id"]] = r   # the last record of each flight has every milestone
        else:
            cards.setdefault(r["entity_id"], []).append(r)

    order = {"kpi.overview": 0, "kpi.checkpoint": 1, "kpi.curbside": 2}
    sheets = []   # (sheet, title, kind, per-minute rows, hourly rows)
    for entity, rows in sorted(cards.items(), key=lambda kv: (order[kv[1][0]["event_type"]], kv[0])):
        kind = rows[0]["event_type"]
        per_minute = [card_row(r) for r in rows]
        hourly = []
        for r in rows:
            if local(r["event_time_epoch_ms"]).minute:
                continue   # the :00 record closes the clock hour before it
            hour = (local(r["event_time_epoch_ms"]) - timedelta(hours=1)).strftime("%H:00")
            if kind == "kpi.checkpoint":
                w = r["hour_avg_wait_minutes"]
                hourly.append({"hour_local": hour, "avg_wait_time_min": w, "threshold_min": r["wait_threshold_minutes"],
                               "over_threshold": None if w is None else ("yes" if w > r["wait_threshold_minutes"] else "no")})
            elif kind == "kpi.curbside":
                hourly.append({"hour_local": hour, "total_vehicles": r["vehicles_this_hour"],
                               **{t: r[f"this_hour_{t}"] for t in ("private_car", "app_cab", "taxi", "bus")},
                               "avg_dwell": mmss(r["avg_dwell_seconds"])})
            else:
                hourly.append({"hour_local": hour, **{k: v for k, v in card_row(r).items()
                                                      if k not in ("time_local", "overview_id", "scope", "name")}})
        title = rows[0].get("desk_name") or rows[0].get("name") or entity
        sheets.append((sheet_name(rows[0]), title, kind, per_minute, hourly))

    def t(v):
        return datetime.fromisoformat(v.replace("Z", "+00:00")).astimezone(TZ).strftime("%Y-%m-%d %H:%M") if v else None
    day0 = day.strftime("%Y-%m-%d")
    flight_rows = []
    for f in sorted(flights.values(), key=lambda f: f["scheduled_time"]):
        arr = f["direction"] == "arrival"
        row = {"flight_id": f["flight_id"], "type": f["direction"], "sector": f["sector"], "status": f["status"],
               "stand_id": f["stand_id"], "gate_id": f["gate_id"],
               "scheduled (STA/STD)": t(f["scheduled_time"]), "estimated (ETA/ETD)": t(f["estimated_time"]),
               "ALDT (landed)": t(f["landing_time"]) if arr else None,
               "AIBT / ATA (in-block)": t(f["in_block_time"]) if arr else None,
               "AOBT / ATD (off-block)": None if arr else t(f["off_block_time"]),
               "ATOT (take-off)": None if arr else t(f["takeoff_time"]),
               "delay_min": f["delay_minutes"], "on_time_+-3min": f["on_time"], "taxi_min": f["taxi_minutes"],
               "pax_forecast": f["pax_forecast"], "pax_actual": f["pax_actual"], "bags": f["bags"]}
        if row["scheduled (STA/STD)"].startswith(day0):
            flight_rows.append(row)

    times = [w.strip() for w in args.snapshot.split(",")] if args.snapshot != "auto" else []
    if not times:
        for card in ("check_in", "passport"):
            desk = [r for rows in cards.values() for r in rows if r.get("card") == card]
            if desk:
                busiest = max(desk, key=lambda r: (r["queue_length"], r["avg_wait_time_seconds"] or 0))
                times.append(local(busiest["event_time_epoch_ms"]).strftime("%H:%M"))
        times = list(dict.fromkeys(times)) or ["08:30"]
    snapshots = {}
    for when in times:
        when = when.strip()
        snap = []
        for sheet, title, _k, rows, _h in sheets:
            r = next((x for x in rows if x["time_local"].endswith(when)), None)
            if r:
                snap.append({"sheet": sheet, "what": title, **{k: v for k, v in r.items() if not k.endswith("_seconds")}})
        snapshots[when.replace(":", "")] = snap

    csv_dir = out / "csv"
    csv_dir.mkdir(exist_ok=True)
    for f in csv_dir.glob("*"):
        f.unlink()
    for sheet, _t, _k, rows, hourly in sheets:
        slug = sheet.lower().replace(" ", "_")
        write_csv(csv_dir / f"{slug}_per_minute.csv", rows)
        write_csv(csv_dir / f"{slug}_hourly.csv", hourly)
    for when, snap in snapshots.items():
        write_csv(csv_dir / f"cards_at_{when}.csv", snap)
    write_csv(csv_dir / "flights.csv", flight_rows)
    # the records exactly as on Kafka, in the same layout as the other sample packs: CSV + JSONL per feed
    raw_dir = out / "kafka_records"
    raw_dir.mkdir(exist_ok=True)
    for f in raw_dir.glob("*"):
        f.unlink()
    tag = day.astimezone(timezone.utc).strftime("%Y%m%dT%H%MZ")
    examples = {}
    for kind in KPI_FEEDS:
        mine = [r for r in records if r["event_type"] == kind]
        if not mine:
            continue
        examples[kind] = next((r for r in mine if r.get("status") == "on_stand" or kind != "flight.status"), mine[0])
        names = list(dict.fromkeys(k for r in mine[:50] for k in columns_for(kind, r)))
        (raw_dir / f"{STEMS[kind]}_{tag}.csv").write_text(to_csv(mine, names), encoding="utf-8")
        with (raw_dir / f"{STEMS[kind]}_{tag}.jsonl").open("w", encoding="utf-8") as fh:
            for r in mine:
                fh.write(json.dumps(r, separators=(",", ":")) + "\n")
    (out / "one_record_per_feed.json").write_text(json.dumps(examples, indent=1) + "\n", encoding="utf-8")

    workbook = out / f"AirportEye_HYD_checkpoint_KPIs_{args.date}.xlsx"
    for old in out.glob("AirportEye_HYD_checkpoint_KPIs_*.xlsx"):
        if old != workbook:
            old.unlink()
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        workbook = None
        print("openpyxl not installed: wrote the CSVs only (pip install -e '.[kpi]')")
    snap_names = ", ".join(f"'Cards at {k}'" for k in snapshots)
    if workbook is not None:
        wb = Workbook()
        bold, head = Font(bold=True), PatternFill("solid", fgColor="DDE6F3")
        wrap = Alignment(wrap_text=True, vertical="top")

        def table(ws, rows):
            keys = list(dict.fromkeys(k for r in rows for k in r))
            ws.append(keys)
            header = ws.max_row
            for cell in ws[header]:
                cell.font, cell.fill = bold, head
            for r in rows:
                ws.append([r.get(k) for k in keys])
            return header

        ws = wb.active
        ws.title = "About"
        for row in [
            ["AirportEye POC: AE-IT widget cards, Hyderabad RGIA Terminal 1"],
            [f"One simulated day: {args.date}, 00:00-24:00 IST. All passengers, flights and vehicles are SYNTHETIC."],
            ["Every value is a field of the live Kafka records (kpi.checkpoint, kpi.curbside, kpi.overview, "
             "flight.status); kafka_records/ has them exactly as sent (CSV and JSONL per topic)."],
            [f"Sheets {snap_names}: every card at that moment (the day's busiest check-in and passport minutes). "
             "Card sheets: formulas, hourly trend, per-minute values."],
            [],
            ["Sheet", "What it covers"],
            *[[s, title] for s, title, *_ in sheets],
            ["Flights", "Every flight of the day: schedule, estimate, ALDT / AIBT / AOBT / ATOT, delay, taxi time, "
                        "forecast and actual passengers, bags (the inputs of the overview cards)."],
            [],
            ["Synthetic per flight", "Schedules, delays, estimates, taxi times, forecasts, bags. In-block / off-block "
                                     "times, passengers and stands come from the simulation."],
            ["Assumed", f"Taxi targets {kpi.taxi_in_target_minutes:g} / {kpi.taxi_out_target_minutes:g} min; trolley "
                        f"stock {kpi.trolleys.total} and use; bags per passenger "
                        f"{kpi.bags_per_passenger.get('domestic')} domestic / "
                        f"{kpi.bags_per_passenger.get('international')} international."],
            ["Happiness Index", f"Wait time against the {kpi.wait_threshold_minutes:g} min threshold until 'FRT' is "
                                "defined."],
            ["Not included", "Bikes (no two-wheelers in the simulation); general aviation flights."],
            ["Generated by", f"scripts/make_checkpoint_kpis.py on {datetime.now(timezone.utc):%Y-%m-%d %H:%M} UTC"],
        ]:
            ws.append(row)
        ws["A1"].font = Font(bold=True, size=14)
        for cell in ws[6]:
            cell.font, cell.fill = bold, head
        ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 30, 110

        for when, snap in snapshots.items():
            ws = wb.create_sheet(f"Cards at {when}")
            table(ws, snap)
            for i in range(1, ws.max_column + 1):
                ws.column_dimensions[get_column_letter(i)].width = 18
        ws = wb.create_sheet("Flights")
        table(ws, flight_rows)
        ws.freeze_panes = "B2"
        for i in range(1, ws.max_column + 1):
            ws.column_dimensions[get_column_letter(i)].width = 17

        for sheet, title, kind, rows, hourly in sheets:
            ws = wb.create_sheet(sheet[:31])
            ws.append([title])
            ws["A1"].font = Font(bold=True, size=13)
            for k, v in formulas(kind, kpi):
                ws.append([k, v])
                ws.cell(ws.max_row, 1).font = bold
                ws.cell(ws.max_row, 2).alignment = wrap
            ws.append([])
            ws.append(["Hourly trend (each row is a whole clock hour)"])
            ws.cell(ws.max_row, 1).font = bold
            table(ws, hourly)
            ws.append([])
            ws.append(["Per minute (live card values)"])
            ws.cell(ws.max_row, 1).font = bold
            header = table(ws, rows)
            ws.freeze_panes = ws.cell(header + 1, 2)
            ws.column_dimensions["A"].width = 26
            ws.column_dimensions["B"].width = 60
            for i in range(3, ws.max_column + 1):
                ws.column_dimensions[get_column_letter(i)].width = 17
        wb.save(workbook)

    (out / "README.md").write_text(
        "# AirportEye POC: AE-IT widget cards, Hyderabad RGIA Terminal 1\n\n"
        f"One simulated day, {args.date} 00:00-24:00 IST. **Synthetic data.** Card values follow the AE-IT KPI sheet "
        "and are the same records as the Kafka KPI topics.\n\n"
        f"- `{workbook.name if workbook else '(no workbook)'}`: About, {snap_names}, Flights, and one sheet per card "
        "(formulas, hourly trend, per-minute values).\n"
        "- `csv/`: the same tables as the workbook.\n"
        "- `kafka_records/`: the records exactly as sent to Kafka, one CSV and one JSONL per topic (one JSONL line = "
        "one message). Load the CSV into ArcGIS with X = `longitude`, Y = `latitude` (WGS84), time = "
        "`event_time_epoch_ms`.\n"
        "- `one_record_per_feed.json`: one example record per topic.\n\n"
        "| Kafka topic | Event type | One record |\n|---|---|---|\n"
        "| `aiop-hyd-topic-kpi-checkpoint-event-gold-v1` | kpi.checkpoint | per desk per minute (track ID `checkpoint_id`) |\n"
        "| `aiop-hyd-topic-kpi-curbside-event-gold-v1` | kpi.curbside | per curbside per minute (`curbside_id`) |\n"
        "| `aiop-hyd-topic-kpi-overview-event-gold-v1` | kpi.overview | airport and terminal per minute (`overview_id`) |\n"
        "| `aiop-hyd-topic-flight-event-silver-v1` | flight.status | per flight at each milestone (`flight_id`) |\n\n"
        "Synthetic per flight: schedules, delays, taxi times, forecasts, bags. Assumed: taxi targets, trolley stock "
        "and use. Not included: bikes, general aviation. Happiness uses wait time against the threshold until 'FRT' "
        "is defined.\n", encoding="utf-8")

    zip_path = Path(args.zip)
    with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as z:
        for f in sorted(out.rglob("*")):
            if f.is_file():
                z.write(f, f"{out.name}/{f.relative_to(out)}")
    for when, snap in snapshots.items():
        print(f"Cards at {when}:", json.dumps({r["sheet"]: {k: r.get(k) for k in (
            "queue_length", "avg_wait_time", "happiness_index", "people_count", "vehicles_this_hour",
            "trolleys_available", "arrival_flights", "departure_flights", "ota_pct", "otd_pct") if r.get(k) is not None}
            for r in snap}, default=str))
    print(f"{len(records):,} records ({len(flight_rows)} flights). Wrote {workbook or csv_dir} and {zip_path} "
          f"({zip_path.stat().st_size / 1e6:.1f} MB)")


if __name__ == "__main__":
    main()
